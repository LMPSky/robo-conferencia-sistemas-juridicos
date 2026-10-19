import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import Workbook, load_workbook  # noqa: E402

from robo_conferencia.listar_pendencias_grupos import (  # noqa: E402
    COLUNA_CNJ_SAIDA,
    COLUNA_STATUS_SAIDA,
    STATUS_CADASTRADO,
    STATUS_PENDENTE,
    STATUS_SEM_NUMERO,
    conferir_grupos_ged,
)


def _criar_grupos(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Nome do grupo", "Data de envio"])
    # Grupo com CNJ que já está na base -> CADASTRADO.
    ws.append(["Proc - 0118284-78.2026.8.04.1000 - Cliente A", "2026-01-05"])
    # Grupo com CNJ que não está na base -> PENDENTE.
    ws.append(["Proc - 0000000-00.2020.8.04.0000 - Cliente B", "2026-01-06"])
    # Grupo sem CNJ identificável no nome -> SEM NÚMERO IDENTIFICADO.
    ws.append(["Lote Triagem Fevereiro - Grupo 12", "2026-02-01"])
    # Grupo com dois CNJs, um cadastrado e outro pendente.
    ws.append(
        [
            (
                "Grupo misto 0118284-78.2026.8.04.1000 e "
                "0667841-11.2025.8.04.1000"
            ),
            "2026-03-01",
        ]
    )
    wb.save(path)


def _criar_legal_one(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Número do Processo", "Cliente"])
    ws.append(["0118284-78.2026.8.04.1000", "Cliente A"])
    wb.save(path)


def test_conferir_grupos_ged(tmp_path: Path):
    grupos_path = tmp_path / "grupos.xlsx"
    legal_one_path = tmp_path / "legal_one.xlsx"
    saida_path = tmp_path / "saida.xlsx"

    _criar_grupos(grupos_path)
    _criar_legal_one(legal_one_path)

    resultado = conferir_grupos_ged(grupos_path, legal_one_path, saida_path)

    assert resultado.processos_legal_one == 1
    assert resultado.grupos_lidos == 4
    assert resultado.grupos_cadastrados == 1
    assert resultado.grupos_pendentes == 2  # linha 2 (PENDENTE) + linha 4 (mista)
    assert resultado.grupos_sem_numero == 1
    assert resultado.linhas_geradas == 5  # 1 + 1 + 1 (sem número) + 2 (grupo misto)

    wb = load_workbook(saida_path)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    cabecalho = linhas[0]
    assert cabecalho[-2:] == (COLUNA_CNJ_SAIDA, COLUNA_STATUS_SAIDA)

    dados = linhas[1:]
    status_por_cnj = {linha[-2]: linha[-1] for linha in dados if linha[-2]}
    assert status_por_cnj["0118284-78.2026.8.04.1000"] == STATUS_CADASTRADO
    assert status_por_cnj["0000000-00.2020.8.04.0000"] == STATUS_PENDENTE
    assert status_por_cnj["0667841-11.2025.8.04.1000"] == STATUS_PENDENTE

    linhas_sem_numero = [linha for linha in dados if linha[-1] == STATUS_SEM_NUMERO]
    assert len(linhas_sem_numero) == 1
    assert linhas_sem_numero[0][0] == "Lote Triagem Fevereiro - Grupo 12"
    assert not linhas_sem_numero[0][-2]


def test_conferir_grupos_ged_planilha_vazia_gera_erro(tmp_path: Path):
    import pytest

    grupos_path = tmp_path / "grupos_vazio.xlsx"
    legal_one_path = tmp_path / "legal_one.xlsx"
    saida_path = tmp_path / "saida.xlsx"

    Workbook().save(grupos_path)
    _criar_legal_one(legal_one_path)

    with pytest.raises(ValueError):
        conferir_grupos_ged(grupos_path, legal_one_path, saida_path)
