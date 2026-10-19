import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import Workbook, load_workbook  # noqa: E402

from robo_conferencia.bater_legal_one import (  # noqa: E402
    COLUNA_CNJ_SAIDA,
    PREENCHIMENTO_MATCH,
    bater_publicacoes_legal_one,
)


def _criar_publicacoes(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Diário", "Data da Publicação", "Conteúdo"])
    ws.append(
        [
            "Diário X",
            "2026-01-01",
            "Processo 0118284-78.2026.8.04.1000 publicado hoje.",
        ]
    )
    ws.append(
        [
            "Diário X",
            "2026-01-02",
            "Processo 0000000-00.2020.8.04.0000 sem match na base.",
        ]
    )
    ws.append(
        [
            "Diário X",
            "2026-01-03",
            (
                "Dois processos: 0118284-78.2026.8.04.1000 e "
                "0667841-11.2025.8.04.1000."
            ),
        ]
    )
    ws.append(
        [
            "Diário X",
            "2026-01-04",
            "Texto sem nenhum número de processo.",
        ]
    )
    wb.save(path)


def _criar_legal_one(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Número do Processo", "Cliente"])
    ws.append(["0118284-78.2026.8.04.1000", "Cliente A"])
    ws.append(["0667841-11.2025.8.04.1000", "Cliente B"])
    wb.save(path)


def test_bater_publicacoes_legal_one(tmp_path: Path):
    pub_path = tmp_path / "publicacoes.xlsx"
    legal_one_path = tmp_path / "legal_one.xlsx"
    saida_path = tmp_path / "saida.xlsx"

    _criar_publicacoes(pub_path)
    _criar_legal_one(legal_one_path)

    resultado = bater_publicacoes_legal_one(pub_path, legal_one_path, saida_path)

    assert resultado.processos_legal_one == 2
    assert resultado.publicacoes_lidas == 4
    # Publicações 1, 2 e 3 citam ao menos um CNJ; a 4 não cita nenhum.
    assert resultado.publicacoes_com_cnj == 3
    # Apenas as publicações 1 e 3 têm CNJ que bate com a base do Legal One.
    assert resultado.publicacoes_com_match == 2
    assert resultado.cnjs_encontrados == 4
    assert resultado.cnjs_com_match == 3
    # Publicação 3 tem 2 CNJs citados -> 2 colunas de CNJ.
    assert resultado.max_cnj_por_publicacao == 2

    wb = load_workbook(saida_path)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    cabecalho = linhas[0]
    assert cabecalho[-2:] == (f"{COLUNA_CNJ_SAIDA} 1", f"{COLUNA_CNJ_SAIDA} 2")
    assert cabecalho[:3] == ("Diário", "Data da Publicação", "Conteúdo")

    # NENHUMA publicação é descartada: as 4 linhas originais permanecem.
    assert len(linhas) - 1 == 4

    por_data = {linha[1]: linha for linha in linhas[1:]}
    # Publicação sem match na base: CNJ continua listado, só não é pintado.
    assert por_data["2026-01-02"][-2:] == ("0000000-00.2020.8.04.0000", None)
    # Publicação sem nenhum CNJ citado: colunas de CNJ vazias.
    assert por_data["2026-01-04"][-2:] == (None, None)
    assert por_data["2026-01-01"][-2:] == ("0118284-78.2026.8.04.1000", None)
    assert por_data["2026-01-03"][-2:] == (
        "0118284-78.2026.8.04.1000",
        "0667841-11.2025.8.04.1000",
    )

    # Verifica a coloração de verde apenas nas células cujo CNJ bate na base.
    cor_por_linha = {}
    for row in ws.iter_rows(min_row=2):
        data = row[1].value
        cores = [cell.fill.fgColor.rgb for cell in row[3:5]]
        cor_por_linha[data] = cores

    verde = PREENCHIMENTO_MATCH.fgColor.rgb
    # 2026-01-01: primeira coluna de CNJ bate -> verde; segunda vazia -> sem cor.
    assert cor_por_linha["2026-01-01"][0] == verde
    assert cor_por_linha["2026-01-01"][1] != verde
    # 2026-01-02: CNJ citado não bate -> sem cor.
    assert cor_por_linha["2026-01-02"][0] != verde
    # 2026-01-03: os dois CNJs batem -> ambas as colunas verdes.
    assert cor_por_linha["2026-01-03"][0] == verde
    assert cor_por_linha["2026-01-03"][1] == verde
    # 2026-01-04: sem CNJ nenhum -> sem cor.
    assert cor_por_linha["2026-01-04"][0] != verde
    assert cor_por_linha["2026-01-04"][1] != verde

