"""Batimento de publicações (Diário Oficial) com a base de processos do Legal One.

Fluxo:
1. Lê a planilha de publicações e extrai, de uma coluna de texto livre (por
   padrão a coluna C), todos os números CNJ encontrados (com ou sem
   pontuação).
2. Lê a planilha de base do Legal One e extrai os números CNJ de uma coluna
   (detectada automaticamente ou informada explicitamente).
3. Gera uma linha de saída para CADA publicação (nenhuma é descartada), com
   todas as colunas originais + uma coluna por CNJ encontrado no texto
   (`Número CNJ`, ou `Número CNJ 1`, `Número CNJ 2`, ... quando há mais de um
   processo citado na mesma publicação). Publicações muito grandes costumam
   citar vários CNJs apenas como referência a outros processos, então o
   robô não tenta adivinhar qual é "o" processo da publicação: ele lista
   todos e deixa a conferência final (via a cor) a cargo de quem for
   trabalhar a planilha.
4. As células de CNJ que existem na base do Legal One são pintadas de verde,
   para destacar rapidamente quais processos citados já estão cadastrados.
"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Sequence, Set

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .cnj import find_cnj_numbers, normalize_cnj

COLUNA_CNJ_SAIDA = "Número CNJ"

# Verde claro (preenchimento) usado para destacar CNJs encontrados na base do
# Legal One.
PREENCHIMENTO_MATCH = PatternFill(
    start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid"
)


@dataclass
class ResultadoBatimento:
    publicacoes_lidas: int
    publicacoes_com_cnj: int
    publicacoes_com_match: int
    cnjs_encontrados: int
    cnjs_com_match: int
    processos_legal_one: int
    max_cnj_por_publicacao: int


def _carregar_planilha(caminho: Path, aba: Optional[str] = None) -> Worksheet:
    if caminho.suffix.lower() == ".xls":
        return _carregar_planilha_xls(caminho, aba)

    wb = load_workbook(caminho, data_only=True, read_only=True)
    if aba:
        if aba not in wb.sheetnames:
            raise ValueError(
                f"A aba '{aba}' não existe em {caminho.name}. "
                f"Abas disponíveis: {', '.join(wb.sheetnames)}"
            )
        return wb[aba]
    # Usa a primeira aba que tenha mais de uma linha (evita abas auxiliares vazias).
    for nome in wb.sheetnames:
        ws = wb[nome]
        if ws.max_row and ws.max_row > 1:
            return ws
    return wb[wb.sheetnames[0]]


def _carregar_planilha_xls(caminho: Path, aba: Optional[str] = None) -> Worksheet:
    """Lê um arquivo .xls (formato antigo, não suportado pelo openpyxl) via
    pandas/xlrd e devolve uma Worksheet openpyxl equivalente em memória, para
    que o restante do código possa tratar .xls e .xlsx da mesma forma."""
    import pandas as pd

    xls = pd.ExcelFile(caminho, engine="xlrd")
    if aba and aba not in xls.sheet_names:
        raise ValueError(
            f"A aba '{aba}' não existe em {caminho.name}. "
            f"Abas disponíveis: {', '.join(xls.sheet_names)}"
        )

    wb = Workbook()
    wb.remove(wb.active)
    escolhida: Optional[Worksheet] = None
    for nome in xls.sheet_names:
        if aba and nome != aba:
            continue
        df = xls.parse(nome, header=None)
        ws = wb.create_sheet(title=nome)
        for linha in df.itertuples(index=False, name=None):
            ws.append([None if pd.isna(valor) else valor for valor in linha])
        if escolhida is None and (aba == nome or (not aba and ws.max_row and ws.max_row > 1)):
            escolhida = ws

    if escolhida is None:
        escolhida = wb[xls.sheet_names[0] if not aba else aba]
    return escolhida


def _resolver_indice_coluna(
    cabecalho: Sequence[object], coluna: str, descricao: str
) -> int:
    """Resolve `coluna` para um índice 0-based de coluna.

    Aceita letra de coluna do Excel (ex: "C"), índice numérico (0-based,
    como string) ou nome do cabeçalho (correspondência exata, sem diferenciar
    maiúsculas/minúsculas).
    """
    coluna = coluna.strip()

    # Letra de coluna do Excel (A, B, C, ..., AA, ...).
    if re.fullmatch(r"[A-Za-z]+", coluna):
        from openpyxl.utils import column_index_from_string

        return column_index_from_string(coluna.upper()) - 1

    # Índice numérico 0-based.
    if coluna.isdigit():
        return int(coluna)

    # Nome do cabeçalho.
    for idx, valor in enumerate(cabecalho):
        if valor is not None and str(valor).strip().lower() == coluna.lower():
            return idx

    raise ValueError(
        f"Não foi possível localizar a coluna '{coluna}' ({descricao}). "
        f"Cabeçalho disponível: {list(cabecalho)}"
    )


def _detectar_coluna_cnj_legal_one(ws: Worksheet) -> int:
    """Detecta automaticamente qual coluna da base do Legal One contém o CNJ.

    Estratégia: para cada coluna, conta quantas das primeiras N linhas contêm
    um número CNJ válido. Escolhe a coluna com maior taxa de acerto.
    """
    linhas = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200), values_only=True))
    if not linhas:
        raise ValueError("A base do Legal One está vazia.")

    num_colunas = max(len(linha) for linha in linhas)
    melhor_idx = -1
    melhor_score = 0
    for idx in range(num_colunas):
        score = 0
        for linha in linhas:
            if idx < len(linha) and linha[idx] is not None:
                if find_cnj_numbers(str(linha[idx])):
                    score += 1
        if score > melhor_score:
            melhor_score = score
            melhor_idx = idx

    if melhor_idx == -1 or melhor_score == 0:
        raise ValueError(
            "Não foi possível detectar automaticamente a coluna de CNJ na base "
            "do Legal One. Informe explicitamente com --coluna-cnj-legal-one."
        )
    return melhor_idx


def _carregar_cnj_legal_one(
    caminho: Path, aba: Optional[str], coluna: Optional[str]
) -> Set[str]:
    ws = _carregar_planilha(caminho, aba)
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return set()

    if coluna:
        idx = _resolver_indice_coluna(linhas[0], coluna, "base do Legal One")
        inicio = 1 if _parece_cabecalho(linhas[0], idx) else 0
    else:
        idx = _detectar_coluna_cnj_legal_one(ws)
        inicio = 1 if _parece_cabecalho(linhas[0], idx) else 0

    processos: Set[str] = set()
    for linha in linhas[inicio:]:
        if idx >= len(linha) or linha[idx] is None:
            continue
        for numero in find_cnj_numbers(str(linha[idx])):
            processos.add(normalize_cnj(numero))
    return processos


def _parece_cabecalho(linha: Sequence[object], idx: int) -> bool:
    if idx >= len(linha) or linha[idx] is None:
        return False
    return not find_cnj_numbers(str(linha[idx]))


def bater_publicacoes_legal_one(
    publicacoes_path: Path,
    legal_one_path: Path,
    saida_path: Path,
    *,
    aba_publicacoes: Optional[str] = None,
    coluna_texto: str = "C",
    aba_legal_one: Optional[str] = None,
    coluna_cnj_legal_one: Optional[str] = None,
) -> ResultadoBatimento:
    """Executa o batimento completo e grava a planilha de saída.

    Nenhuma publicação é descartada: todas as linhas da planilha original são
    mantidas na saída, com uma coluna por CNJ encontrado no texto. As células
    cujo CNJ existe na base do Legal One são pintadas de verde.

    Retorna um `ResultadoBatimento` com estatísticas da execução.
    """
    processos_legal_one = _carregar_cnj_legal_one(
        legal_one_path, aba_legal_one, coluna_cnj_legal_one
    )

    ws_pub = _carregar_planilha(publicacoes_path, aba_publicacoes)
    linhas = list(ws_pub.iter_rows(values_only=True))
    if not linhas:
        raise ValueError("A planilha de publicações está vazia.")

    cabecalho = list(linhas[0])
    idx_texto = _resolver_indice_coluna(cabecalho, coluna_texto, "planilha de publicações")
    linhas_dados = linhas[1:]

    # Primeira passada: para cada publicação (sem descartar nenhuma), extrai
    # todos os CNJs citados no texto, na ordem em que aparecem.
    cnjs_por_linha: List[List[str]] = []
    max_cnj = 0
    publicacoes_com_cnj = 0
    cnjs_encontrados = 0
    cnjs_com_match = 0
    for linha in linhas_dados:
        texto = linha[idx_texto] if idx_texto < len(linha) else None
        numeros = find_cnj_numbers(str(texto)) if texto else []
        cnjs_por_linha.append(numeros)
        if numeros:
            publicacoes_com_cnj += 1
        max_cnj = max(max_cnj, len(numeros))
        cnjs_encontrados += len(numeros)
        cnjs_com_match += sum(
            1 for n in numeros if normalize_cnj(n) in processos_legal_one
        )

    wb_saida = Workbook()
    ws_saida = wb_saida.active
    ws_saida.title = "Publicações"
    colunas_cnj = _nomes_colunas_cnj(max_cnj)
    ws_saida.append([*cabecalho, *colunas_cnj])

    idx_primeira_coluna_cnj = len(cabecalho)
    publicacoes_com_match = 0
    for linha, numeros in zip(linhas_dados, cnjs_por_linha):
        preenchimento = list(numeros) + [None] * (max_cnj - len(numeros))
        ws_saida.append([*linha, *preenchimento])
        tem_match_na_linha = False
        linha_excel = ws_saida.max_row
        for offset, numero in enumerate(numeros):
            if normalize_cnj(numero) in processos_legal_one:
                tem_match_na_linha = True
                coluna_excel = idx_primeira_coluna_cnj + offset + 1
                ws_saida.cell(row=linha_excel, column=coluna_excel).fill = (
                    PREENCHIMENTO_MATCH
                )
        if tem_match_na_linha:
            publicacoes_com_match += 1

    saida_path.parent.mkdir(parents=True, exist_ok=True)
    wb_saida.save(saida_path)

    return ResultadoBatimento(
        publicacoes_lidas=len(linhas_dados),
        publicacoes_com_cnj=publicacoes_com_cnj,
        publicacoes_com_match=publicacoes_com_match,
        cnjs_encontrados=cnjs_encontrados,
        cnjs_com_match=cnjs_com_match,
        processos_legal_one=len(processos_legal_one),
        max_cnj_por_publicacao=max_cnj,
    )


def _nomes_colunas_cnj(quantidade: int) -> List[str]:
    """Gera os nomes das colunas de CNJ na saída.

    Com apenas 1 coluna necessária, usa o nome simples (`Número CNJ`). Com
    mais de uma, numera cada coluna (`Número CNJ 1`, `Número CNJ 2`, ...) para
    não perder nenhum processo citado na mesma publicação.
    """
    if quantidade <= 1:
        return [COLUNA_CNJ_SAIDA]
    return [f"{COLUNA_CNJ_SAIDA} {i}" for i in range(1, quantidade + 1)]


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Extrai todos os números CNJ citados em cada publicação, bate "
            "com a base do Legal One e gera uma planilha com todas as "
            "publicações, destacando em verde os CNJs encontrados na base."
        )
    )
    parser.add_argument(
        "publicacoes", type=Path, help="Planilha de publicações (.xlsx ou .xls)"
    )
    parser.add_argument(
        "legal_one", type=Path, help="Planilha de base do Legal One (.xlsx ou .xls)"
    )
    parser.add_argument("saida", type=Path, help="Caminho da planilha de saída (.xlsx)")
    parser.add_argument(
        "--aba-publicacoes", default=None, help="Nome da aba da planilha de publicações"
    )
    parser.add_argument(
        "--coluna-texto",
        default="C",
        help=(
            "Coluna com o texto livre a analisar na planilha de publicações. "
            "Aceita letra do Excel (ex: C), índice 0-based ou nome do cabeçalho. "
            "Padrão: C"
        ),
    )
    parser.add_argument(
        "--aba-legal-one", default=None, help="Nome da aba da planilha de base do Legal One"
    )
    parser.add_argument(
        "--coluna-cnj-legal-one",
        default=None,
        help=(
            "Coluna com o número CNJ na base do Legal One. Se omitida, é "
            "detectada automaticamente."
        ),
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = _build_arg_parser()
    args = parser.parse_args(argv)

    resultado = bater_publicacoes_legal_one(
        args.publicacoes,
        args.legal_one,
        args.saida,
        aba_publicacoes=args.aba_publicacoes,
        coluna_texto=args.coluna_texto,
        aba_legal_one=args.aba_legal_one,
        coluna_cnj_legal_one=args.coluna_cnj_legal_one,
    )

    print(f"Processos na base do Legal One: {resultado.processos_legal_one}")
    print(f"Publicações lidas: {resultado.publicacoes_lidas}")
    print(f"Publicações com ao menos um CNJ citado: {resultado.publicacoes_com_cnj}")
    print(f"Publicações com ao menos um CNJ encontrado na base: {resultado.publicacoes_com_match}")
    print(f"Total de CNJs citados: {resultado.cnjs_encontrados}")
    print(f"Total de CNJs encontrados na base (pintados de verde): {resultado.cnjs_com_match}")
    print(f"Máximo de CNJs citados em uma mesma publicação: {resultado.max_cnj_por_publicacao}")
    print(f"Planilha gerada em: {args.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
