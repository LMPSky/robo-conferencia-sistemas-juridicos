"""Batimento entre planilhas de cadastro de processos (baixadas do GED) e a
base de processos do Legal One.

Contexto de negócio
--------------------
Diariamente, times diferentes (Advogado, Robô TI, Sistema Cliente, Triagem
Publicações) sobem no GED do Legal One planilhas "Planilha de Cadastro -
<tipo>" contendo processos que precisam ser cadastrados como novos processos
no Legal One. Este módulo verifica se cada processo listado nessas planilhas
já foi efetivamente cadastrado no Legal One e gera um relatório de pendências.

Fluxo geral
-----------
1. Carrega a base de processos do Legal One (``entrada/base_legal_one``) e
   constrói um conjunto de "chaves conhecidas" normalizadas a partir das
   colunas Número de CNJ, Número antigo, Outro número e Pasta.
2. Para cada planilha de cadastro individual (``entrada/planilhas_cadastro``),
   localiza a aba de cadastro e as colunas relevantes por heurística
   (regex tolerante a variações de nome/acentuação), e verifica se o número
   de cada processo já está cadastrado no Legal One.
3. Opcionalmente enriquece cada linha com tipo/data de envio a partir do
   export de metadados do GED (``entrada/export_ged``).
4. Gera um relatório ``saida/Batimento_AAAAMMDD_HHMMSS.xlsx`` com as abas
   "Resumo", "Pendencias" e "Detalhe completo".
"""
from __future__ import annotations

import difflib
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Constantes
# --------------------------------------------------------------------------

STATUS_CADASTRADO = "CADASTRADO"
STATUS_PENDENTE = "PENDENTE"
STATUS_SEM_NUMERO = "SEM NÚMERO IDENTIFICADO"
TIPO_NAO_IDENTIFICADO = "Não identificado"

PASTA_PREFIXOS = (
    "proc -",
    "proc-",
    "proc",
)

_FILL_PENDENTE = PatternFill(start_color="FFF8CBCB", end_color="FFF8CBCB", fill_type="solid")
_FILL_SEM_NUMERO = PatternFill(start_color="FFFFF2B2", end_color="FFFFF2B2", fill_type="solid")
_FILL_CADASTRADO = PatternFill(start_color="FFC6E8C6", end_color="FFC6E8C6", fill_type="solid")
_FILL_HEADER = PatternFill(start_color="FF2F3640", end_color="FF2F3640", fill_type="solid")
_FONT_HEADER = Font(color="FFFFFFFF", bold=True)


# --------------------------------------------------------------------------
# Normalização de texto / números
# --------------------------------------------------------------------------

def remover_acentos(texto: str) -> str:
    """Remove acentuação de uma string, preservando os demais caracteres."""
    if texto is None:
        return ""
    nfkd = unicodedata.normalize("NFKD", str(texto))
    return "".join(ch for ch in nfkd if not unicodedata.combining(ch))


def normalizar_cabecalho(texto: object) -> str:
    """Normaliza um cabeçalho de coluna para comparação por regex: minúsculo,
    sem acentos, espaços colapsados."""
    if texto is None:
        return ""
    limpo = remover_acentos(str(texto)).lower().strip()
    return re.sub(r"\s+", " ", limpo)


def normalizar_numero(valor: object) -> str:
    """Normaliza um número de processo: remove espaços nas pontas, coloca em
    maiúsculas e colapsa espaços internos duplicados."""
    if valor is None:
        return ""
    texto = str(valor).strip()
    if not texto or texto.lower() == "nan":
        return ""
    texto = re.sub(r"\s+", " ", texto).upper()
    return texto


def normalizar_pasta(valor: object) -> str:
    """Normaliza o valor da coluna 'Pasta' removendo o prefixo 'Proc - '."""
    base = normalizar_numero(valor)
    base_lower = base.lower()
    for prefixo in PASTA_PREFIXOS:
        if base_lower.startswith(prefixo):
            resto = base[len(prefixo):].strip(" -")
            return resto
    return base


def gerar_variantes(valor: object, *, eh_pasta: bool = False) -> Set[str]:
    """Gera variantes normalizadas de um número de processo/pasta para
    aumentar a chance de match entre bases com formatação diferente.

    Variantes geradas: valor normalizado, sem prefixo 'Proc - ' (se pasta),
    sem barras, e apenas dígitos (para casar formatos CNJ com/sem pontuação).
    """
    variantes: Set[str] = set()
    base = normalizar_pasta(valor) if eh_pasta else normalizar_numero(valor)
    if not base:
        return variantes

    variantes.add(base)
    sem_barra = base.replace("/", "")
    variantes.add(sem_barra)

    apenas_digitos = re.sub(r"\D", "", base)
    if apenas_digitos:
        variantes.add(apenas_digitos)

    variantes.discard("")
    return variantes


# --------------------------------------------------------------------------
# Detecção de colunas por heurística (regex tolerante a acentos/sinônimos)
# --------------------------------------------------------------------------

def encontrar_coluna(colunas: Sequence[object], termos_obrigatorios: Sequence[str]) -> Optional[object]:
    """Retorna o nome original da primeira coluna cujo cabeçalho normalizado
    contém todos os termos informados (cada termo é um regex simples)."""
    for coluna in colunas:
        normalizado = normalizar_cabecalho(coluna)
        if all(re.search(termo, normalizado) for termo in termos_obrigatorios):
            return coluna
    return None


def encontrar_coluna_com_alternativas(
    colunas: Sequence[object], alternativas: Sequence[Sequence[str]]
) -> Optional[object]:
    """Como `encontrar_coluna`, mas aceita múltiplas alternativas de termos
    para o mesmo campo: tenta a primeira alternativa (em todas as colunas);
    se nenhuma coluna bater, tenta a próxima alternativa, e assim por diante.

    Isso permite que um campo interno (ex: ``numero_cnj``) reconheça vários
    "sinônimos" de cabeçalho (ex: planilhas de times/robôs diferentes),
    mantendo a prioridade da alternativa mais específica/consolidada.
    """
    for termos_obrigatorios in alternativas:
        coluna = encontrar_coluna(colunas, termos_obrigatorios)
        if coluna is not None:
            return coluna
    return None


# Cada campo interno mapeia para uma tupla de alternativas de termos; cada
# alternativa é, por sua vez, uma tupla de regexes que devem casar todos
# (mesmo formato aceito por `encontrar_coluna`). A primeira alternativa que
# encontrar uma coluna correspondente é usada.
COLUNAS_CADASTRO = {
    "numero_cnj": (
        (r"cnj",),
        # "Número do processo" (ex: planilhas do Robô Protocolo), mas sem
        # roubar colunas de "outro número" (ex: "Outro Número do Processo").
        (r"^(?!.*outro)", r"numero", r"processo"),
        # Fallback mais genérico (prioridade mais baixa): planilhas de 2
        # colunas (ex: "Processos - Estratégico.xlsx"/"Processos -
        # Tributário.xlsx") com cabeçalho simplesmente "Processos:", sem a
        # palavra "número" junto. Mesma proteção contra "outro número".
        (r"^(?!.*outro)", r"\bprocessos?\b"),
    ),
    "outro_numero": ((r"outro", r"numer"),),
    "cliente_principal": (
        (r"client", r"principal"),
        (r"\bpartes\b",),
    ),
    "responsavel_principal": (
        (r"respons", r"principal"),
        (r"advogado", r"respons"),
    ),
    "status": ((r"status",),),
    # Remetente/caixa relacionada ao e-mail que originou o registro (quando
    # disponível na planilha de origem, ex: planilhas do Robô Protocolo).
    # Nem sempre há uma coluna de remetente explícita, então caímos para a
    # caixa monitorada e, por último, o e-mail do destinatário como melhor
    # aproximação disponível.
    "remetente_email": (
        (r"remetente",),
        (r"de quem enviou",),
        (r"caixa", r"monitorada"),
        (r"e-?mail", r"destinatari"),
    ),
    # Data de envio do e-mail que originou o registro (não confundir com a
    # coluna "Data de envio" existente, que vem do export_ged e representa a
    # data de anexação/geração no GED).
    "data_envio_email": (
        (r"data", r"hora", r"e-?mail"),
        (r"data", r"envio", r"e-?mail"),
        (r"data", r"e-?mail"),
    ),
    # Data do processo/distribuição (ex: planilhas de 2 colunas "Processos:"
    # / "Data:"). Termo genérico demais para arriscar confundir com a data
    # do e-mail, por isso exige que o cabeçalho não mencione "e-mail"/"email".
    "data_processo": (
        (r"^(?!.*e-?mail)", r"\bdata\b"),
    ),
}


# --------------------------------------------------------------------------
# Pastas de entrada/saída
# --------------------------------------------------------------------------

@dataclass
class EstruturaPastas:
    entrada: Path
    base_legal_one: Path
    planilhas_cadastro: Path
    export_ged: Path
    saida: Path


def garantir_estrutura_pastas(base_dir: Path) -> EstruturaPastas:
    """Cria a estrutura de pastas esperada (se ainda não existir), com um
    arquivo `.gitkeep` em cada pasta vazia para preservar a estrutura no git."""
    entrada = base_dir / "entrada"
    estrutura = EstruturaPastas(
        entrada=entrada,
        base_legal_one=entrada / "base_legal_one",
        planilhas_cadastro=entrada / "planilhas_cadastro",
        export_ged=entrada / "export_ged",
        saida=base_dir / "saida",
    )
    for pasta in (
        estrutura.base_legal_one,
        estrutura.planilhas_cadastro,
        estrutura.export_ged,
        estrutura.saida,
    ):
        pasta.mkdir(parents=True, exist_ok=True)
        gitkeep = pasta / ".gitkeep"
        if not any(pasta.iterdir()):
            gitkeep.touch()
    return estrutura


def _listar_arquivos_excel(pasta: Path, extensoes: Sequence[str]) -> List[Path]:
    arquivos = [
        p
        for p in sorted(pasta.iterdir())
        if p.is_file() and p.suffix.lower() in extensoes and not p.name.startswith("~$")
    ]
    return arquivos


# --------------------------------------------------------------------------
# Base do Legal One
# --------------------------------------------------------------------------

def carregar_base_legal_one(pasta_base: Path) -> pd.DataFrame:
    """Carrega o(s) export(s) da base do Legal One.

    A primeira linha do arquivo é um título solto (ex: 'usuario@empresa.com.br
    - 01/01/2026') e deve ser ignorada; a segunda linha tem os cabeçalhos
    reais, por isso usamos `skiprows=1` no `read_excel`.
    """
    arquivos = _listar_arquivos_excel(pasta_base, (".xlsx", ".xls"))
    if not arquivos:
        raise FileNotFoundError(
            "Nenhum arquivo encontrado em 'entrada/base_legal_one'. "
            "Coloque ali o export (.xlsx) da base de processos do Legal One "
            "antes de rodar o robô."
        )

    dataframes = []
    for arquivo in arquivos:
        try:
            df = pd.read_excel(arquivo, skiprows=1, dtype=str)
        except Exception:
            # Fallback: talvez não haja a linha de título solta neste export.
            df = pd.read_excel(arquivo, dtype=str)
        dataframes.append(df)

    return pd.concat(dataframes, ignore_index=True, sort=False)


def construir_chaves_conhecidas(df_legal_one: pd.DataFrame) -> Set[str]:
    """Constrói o conjunto de chaves normalizadas conhecidas pelo Legal One,
    a partir das colunas Número de CNJ, Número antigo, Outro número e Pasta."""
    chaves: Set[str] = set()

    coluna_pasta = encontrar_coluna(df_legal_one.columns, (r"pasta",))
    coluna_cnj = encontrar_coluna(df_legal_one.columns, (r"cnj",))
    coluna_numero_antigo = encontrar_coluna(df_legal_one.columns, (r"numero antigo|numero.?antigo",)) \
        or encontrar_coluna(df_legal_one.columns, (r"antigo",))
    coluna_outro_numero = encontrar_coluna(df_legal_one.columns, (r"outro", r"numer"))

    for coluna, eh_pasta in (
        (coluna_cnj, False),
        (coluna_numero_antigo, False),
        (coluna_outro_numero, False),
        (coluna_pasta, True),
    ):
        if coluna is None or coluna not in df_legal_one.columns:
            continue
        for valor in df_legal_one[coluna].dropna().unique():
            chaves.update(gerar_variantes(valor, eh_pasta=eh_pasta))

    chaves.discard("")
    return chaves


# --------------------------------------------------------------------------
# Planilhas de cadastro
# --------------------------------------------------------------------------

def localizar_aba_cadastro(caminho_arquivo: Path) -> str:
    """Localiza o nome da aba de cadastro dentro da planilha: procura por
    aba chamada (aproximadamente) 'CADASTRO'; se não achar, procura a
    primeira aba cujo cabeçalho contenha uma coluna com 'CNJ'; senão usa a
    primeira aba disponível."""
    xls = pd.ExcelFile(caminho_arquivo)
    nomes_normalizados = {aba: normalizar_cabecalho(aba) for aba in xls.sheet_names}

    for aba, nome_normalizado in nomes_normalizados.items():
        if "cadastro" in nome_normalizado:
            return aba

    for aba in xls.sheet_names:
        try:
            cabecalho = pd.read_excel(xls, sheet_name=aba, nrows=0).columns
        except Exception:
            continue
        if encontrar_coluna(cabecalho, (r"cnj",)) is not None:
            return aba

    # Fallback adicional: planilhas de aba única (sem nenhuma coluna "CNJ"),
    # como o export do "Robô Protocolo" (coluna "Número do processo") ou
    # planilhas simples de 2 colunas (coluna "Processos:"), identificadas
    # reaproveitando as mesmas alternativas de `numero_cnj`.
    for aba in xls.sheet_names:
        try:
            cabecalho = pd.read_excel(xls, sheet_name=aba, nrows=0).columns
        except Exception:
            continue
        if encontrar_coluna_com_alternativas(cabecalho, COLUNAS_CADASTRO["numero_cnj"]) is not None:
            return aba

    return xls.sheet_names[0]


def _ler_coluna_como_datetime(
    caminho_arquivo: Path, aba: str, coluna: Optional[str]
) -> Optional[pd.Series]:
    """Relê uma única coluna da planilha sem forçar `dtype=str`, convertendo
    para datetime. Usado para colunas de data (`data_envio_email`,
    `data_processo`) que, se lidas como texto, perderiam a formatação
    original (ex: viram "2026-10-08 00:00:00" em vez de uma data real)."""
    if coluna is None:
        return None
    try:
        df_coluna = pd.read_excel(caminho_arquivo, sheet_name=aba, usecols=[coluna])
        return pd.to_datetime(df_coluna[coluna], errors="coerce")
    except Exception:
        return None


def carregar_planilha_cadastro(caminho_arquivo: Path) -> pd.DataFrame:
    """Carrega a aba de cadastro de uma planilha individual e renomeia as
    colunas relevantes para nomes padronizados internos."""
    aba = localizar_aba_cadastro(caminho_arquivo)
    df = pd.read_excel(caminho_arquivo, sheet_name=aba, dtype=str)

    mapeamento: Dict[str, str] = {}
    for nome_interno, alternativas in COLUNAS_CADASTRO.items():
        coluna = encontrar_coluna_com_alternativas(df.columns, alternativas)
        if coluna is not None:
            mapeamento[coluna] = nome_interno

    # As colunas de data (ex: "Data/Hora do e-mail" no formato Robô
    # Protocolo, ou "Data:" nas planilhas simples de 2 colunas) são valores
    # datetime reais no Excel. Forçar `dtype=str` na leitura acima
    # transformaria essas datas em texto ISO cru (ex: "2026-10-08
    # 00:00:00"), então relemos cada uma delas sem forçar o tipo e
    # convertemos explicitamente para datetime, preservando as demais
    # colunas (numero_cnj, outro_numero, cliente_principal,
    # responsavel_principal, status, remetente_email) como texto.
    series_por_coluna_data: Dict[str, pd.Series] = {}
    for nome_interno_data in ("data_envio_email", "data_processo"):
        coluna_origem = next(
            (origem for origem, destino in mapeamento.items() if destino == nome_interno_data),
            None,
        )
        serie = _ler_coluna_como_datetime(caminho_arquivo, aba, coluna_origem)
        if serie is not None:
            series_por_coluna_data[nome_interno_data] = serie

    df = df.rename(columns=mapeamento)

    for nome_interno in COLUNAS_CADASTRO:
        if nome_interno not in df.columns:
            df[nome_interno] = None

    for nome_interno_data, serie in series_por_coluna_data.items():
        df[nome_interno_data] = serie.values

    # Ignora linhas totalmente vazias nas colunas-chave.
    colunas_chave = ["numero_cnj", "outro_numero", "cliente_principal", "responsavel_principal"]
    mascara_vazia = df[colunas_chave].apply(
        lambda col: col.map(lambda v: normalizar_numero(v) == "")
    ).all(axis=1)
    df = df.loc[~mascara_vazia].reset_index(drop=True)

    return df


# --------------------------------------------------------------------------
# Export de metadados do GED (opcional)
# --------------------------------------------------------------------------

def normalizar_nome_arquivo(nome: object) -> str:
    """Normaliza um nome de arquivo/planilha para comparação por
    aproximação: minúsculo, sem acentos, sem extensão, espaços colapsados."""
    if nome is None:
        return ""
    texto = remover_acentos(str(nome)).lower().strip()
    texto = re.sub(r"\.(xlsx|xls|xlsm)$", "", texto)
    texto = re.sub(r"[_\-]+", " ", texto)
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


def carregar_export_ged(pasta_export: Path) -> Optional[pd.DataFrame]:
    """Carrega o export de metadados do GED, se houver algum arquivo na
    pasta. Retorna None se a pasta estiver vazia (é opcional)."""
    arquivos = _listar_arquivos_excel(pasta_export, (".xlsx", ".xls"))
    if not arquivos:
        return None

    dataframes = []
    for arquivo in arquivos:
        engine = "xlrd" if arquivo.suffix.lower() == ".xls" else "openpyxl"
        try:
            df = pd.read_excel(arquivo, engine=engine, dtype=str)
        except Exception:
            df = pd.read_excel(arquivo, dtype=str)
        dataframes.append(df)

    df_final = pd.concat(dataframes, ignore_index=True, sort=False)

    coluna_data = encontrar_coluna(df_final.columns, (r"anexad|gerad",))
    coluna_nome = encontrar_coluna(df_final.columns, (r"descri|nome",))
    coluna_tipo = encontrar_coluna(df_final.columns, (r"tipo",))

    df_final.attrs["coluna_data"] = coluna_data
    df_final.attrs["coluna_nome"] = coluna_nome
    df_final.attrs["coluna_tipo"] = coluna_tipo
    return df_final


def encontrar_tipo_e_data(
    nome_arquivo: str, df_ged: Optional[pd.DataFrame]
) -> Tuple[str, str]:
    """Encontra o tipo de planilha e a data de envio correspondentes a um
    arquivo de cadastro, por aproximação de nome com o export do GED."""
    if df_ged is None or df_ged.empty:
        return TIPO_NAO_IDENTIFICADO, ""

    coluna_nome = df_ged.attrs.get("coluna_nome")
    coluna_tipo = df_ged.attrs.get("coluna_tipo")
    coluna_data = df_ged.attrs.get("coluna_data")
    if coluna_nome is None:
        return TIPO_NAO_IDENTIFICADO, ""

    alvo = normalizar_nome_arquivo(nome_arquivo)
    if not alvo:
        return TIPO_NAO_IDENTIFICADO, ""

    melhor_score = 0.0
    melhor_idx = None
    for idx, valor in df_ged[coluna_nome].items():
        candidato = normalizar_nome_arquivo(valor)
        if not candidato:
            continue
        if alvo in candidato or candidato in alvo:
            score = max(len(alvo), len(candidato)) and (
                min(len(alvo), len(candidato)) / max(len(alvo), len(candidato))
            )
            score = 0.9 + score * 0.1  # containment sempre priorizado
        else:
            score = difflib.SequenceMatcher(None, alvo, candidato).ratio()
        if score > melhor_score:
            melhor_score = score
            melhor_idx = idx

    if melhor_idx is None or melhor_score < 0.5:
        return TIPO_NAO_IDENTIFICADO, ""

    tipo = TIPO_NAO_IDENTIFICADO
    if coluna_tipo is not None:
        valor_tipo = df_ged.at[melhor_idx, coluna_tipo]
        if valor_tipo and str(valor_tipo).strip().lower() != "nan":
            tipo = str(valor_tipo).strip()

    data = ""
    if coluna_data is not None:
        valor_data = df_ged.at[melhor_idx, coluna_data]
        if valor_data and str(valor_data).strip().lower() != "nan":
            data = str(valor_data).strip()

    return tipo, data


# --------------------------------------------------------------------------
# Lógica central de batimento
# --------------------------------------------------------------------------

def determinar_numero_busca(linha: pd.Series) -> Optional[str]:
    cnj = normalizar_numero(linha.get("numero_cnj"))
    if cnj:
        return cnj
    outro = normalizar_numero(linha.get("outro_numero"))
    if outro:
        return outro
    return None


def valor_texto_ou_vazio(valor: object) -> str:
    """Converte um valor opcional de uma célula para texto, tratando ausência
    como string vazia.

    Necessário porque `DataFrame.iterrows()` pode converter um `None` (usado
    para preencher colunas internas ausentes na planilha de origem, ex:
    `remetente_email`/`data_envio_email` quando a planilha não tem essa
    informação) em `NaN` (float) ao montar a `Series` da linha, e `NaN` é
    "truthy" em Python, então um simples `valor or ""` não seria suficiente.
    """
    if valor is None:
        return ""
    if isinstance(valor, float) and pd.isna(valor):
        return ""
    texto = str(valor).strip()
    if texto.lower() == "nan":
        return ""
    return texto


def formatar_data(valor: object) -> str:
    """Formata uma data (ex: data de envio do e-mail, data do processo)
    como "dd/mm/AAAA" (ou "dd/mm/AAAA HH:MM" quando houver um horário
    diferente de meia-noite genuinamente presente na planilha de origem).

    `valor` normalmente já é um `Timestamp`/`datetime` (ver
    `carregar_planilha_cadastro`, que preserva o tipo datetime nativo do
    Excel para essas colunas). Se a conversão falhar ou o valor vier vazio,
    retorna string vazia, sem erro.
    """
    if valor is None:
        return ""
    if isinstance(valor, float) and pd.isna(valor):
        return ""

    data = valor if isinstance(valor, (pd.Timestamp, datetime)) else pd.to_datetime(valor, errors="coerce")
    if data is None or pd.isna(data):
        return ""

    if data.hour == 0 and data.minute == 0 and data.second == 0:
        return data.strftime("%d/%m/%Y")
    return data.strftime("%d/%m/%Y %H:%M")


def verificar_status(numero_busca: Optional[str], chaves_conhecidas: Set[str]) -> str:
    if numero_busca is None:
        return STATUS_SEM_NUMERO
    variantes = gerar_variantes(numero_busca)
    if variantes & chaves_conhecidas:
        return STATUS_CADASTRADO
    return STATUS_PENDENTE


def processar_planilhas_cadastro(
    pasta_planilhas: Path,
    chaves_conhecidas: Set[str],
    df_ged: Optional[pd.DataFrame],
) -> pd.DataFrame:
    arquivos = _listar_arquivos_excel(pasta_planilhas, (".xlsx", ".xls"))
    if not arquivos:
        raise FileNotFoundError(
            "Nenhum arquivo encontrado em 'entrada/planilhas_cadastro'. "
            "Coloque ali as planilhas 'PLANILHA DE CADASTRO - ...' baixadas "
            "do GED antes de rodar o robô."
        )

    linhas_resultado: List[Dict[str, object]] = []
    for arquivo in arquivos:
        tipo, data_envio = encontrar_tipo_e_data(arquivo.name, df_ged)
        df_cadastro = carregar_planilha_cadastro(arquivo)

        for _, linha in df_cadastro.iterrows():
            numero_busca = determinar_numero_busca(linha)
            status = verificar_status(numero_busca, chaves_conhecidas)

            if status == STATUS_PENDENTE:
                observacao = "Processo não localizado na base do Legal One."
            elif status == STATUS_SEM_NUMERO:
                observacao = "Linha sem Número CNJ e sem Outro número preenchidos; checagem manual."
            else:
                observacao = ""

            linhas_resultado.append(
                {
                    "Arquivo de origem": arquivo.name,
                    "Tipo": tipo,
                    "Cliente principal": valor_texto_ou_vazio(linha.get("cliente_principal")),
                    "Responsável principal": valor_texto_ou_vazio(linha.get("responsavel_principal")),
                    "Remetente do e-mail": valor_texto_ou_vazio(linha.get("remetente_email")),
                    "Data de envio do e-mail": formatar_data(linha.get("data_envio_email")),
                    "Data do processo": formatar_data(linha.get("data_processo")),
                    "Status planilha": valor_texto_ou_vazio(linha.get("status")),
                    "Número CNJ": valor_texto_ou_vazio(linha.get("numero_cnj")),
                    "Outro número": valor_texto_ou_vazio(linha.get("outro_numero")),
                    "Número buscado": numero_busca or "",
                    "Status": status,
                    "Data de envio": data_envio,
                    "Observação": observacao,
                }
            )

    if not linhas_resultado:
        raise ValueError(
            "Nenhuma linha válida encontrada nas planilhas de cadastro "
            "(todas as linhas estavam vazias nas colunas-chave)."
        )

    return pd.DataFrame(linhas_resultado)


def gerar_resumo(df_detalhe: pd.DataFrame) -> pd.DataFrame:
    agrupado = (
        df_detalhe.groupby(["Arquivo de origem", "Tipo"], dropna=False)["Status"]
        .value_counts()
        .unstack(fill_value=0)
    )
    for status in (STATUS_CADASTRADO, STATUS_PENDENTE, STATUS_SEM_NUMERO):
        if status not in agrupado.columns:
            agrupado[status] = 0

    agrupado["Total de processos"] = agrupado[
        [STATUS_CADASTRADO, STATUS_PENDENTE, STATUS_SEM_NUMERO]
    ].sum(axis=1)
    agrupado["% cadastrado"] = (
        agrupado[STATUS_CADASTRADO] / agrupado["Total de processos"].replace(0, pd.NA) * 100
    ).round(1)

    agrupado = agrupado.reset_index().rename(
        columns={
            STATUS_CADASTRADO: "Cadastrados",
            STATUS_PENDENTE: "Pendentes",
            STATUS_SEM_NUMERO: "Sem número identificado",
        }
    )

    colunas_ordenadas = [
        "Arquivo de origem",
        "Tipo",
        "Total de processos",
        "Cadastrados",
        "Pendentes",
        "Sem número identificado",
        "% cadastrado",
    ]
    return agrupado[colunas_ordenadas]


# --------------------------------------------------------------------------
# Geração do relatório Excel
# --------------------------------------------------------------------------

def _autoajustar_largura_colunas(ws) -> None:
    for idx, coluna in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        maior = max((len(str(c.value)) for c in coluna if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(maior + 2, 12), 60)


def _escrever_aba(writer: pd.ExcelWriter, df: pd.DataFrame, nome_aba: str) -> None:
    df.to_excel(writer, sheet_name=nome_aba, index=False)
    ws = writer.sheets[nome_aba]
    for celula in ws[1]:
        celula.fill = _FILL_HEADER
        celula.font = _FONT_HEADER
        celula.alignment = Alignment(horizontal="center")
    _autoajustar_largura_colunas(ws)


def _aplicar_formatacao_condicional(ws, colunas: Sequence[str]) -> None:
    if "Status" not in colunas:
        return
    idx_status = list(colunas).index("Status") + 1
    mapa_fill = {
        STATUS_PENDENTE: _FILL_PENDENTE,
        STATUS_SEM_NUMERO: _FILL_SEM_NUMERO,
        STATUS_CADASTRADO: _FILL_CADASTRADO,
    }
    for linha in ws.iter_rows(min_row=2, max_row=ws.max_row):
        valor_status = linha[idx_status - 1].value
        fill = mapa_fill.get(valor_status)
        if fill is not None:
            for celula in linha:
                celula.fill = fill


def gerar_relatorio(
    df_detalhe: pd.DataFrame, df_resumo: pd.DataFrame, caminho_saida: Path
) -> None:
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)

    df_pendencias = df_detalhe.loc[
        df_detalhe["Status"].isin([STATUS_PENDENTE, STATUS_SEM_NUMERO]),
        [
            "Arquivo de origem",
            "Tipo",
            "Responsável principal",
            "Cliente principal",
            "Remetente do e-mail",
            "Data de envio do e-mail",
            "Data do processo",
            "Número buscado",
            "Status",
            "Observação",
        ],
    ].reset_index(drop=True)

    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        _escrever_aba(writer, df_resumo, "Resumo")
        _escrever_aba(writer, df_pendencias, "Pendencias")
        _escrever_aba(writer, df_detalhe, "Detalhe completo")

        _aplicar_formatacao_condicional(writer.sheets["Pendencias"], df_pendencias.columns)
        _aplicar_formatacao_condicional(writer.sheets["Detalhe completo"], df_detalhe.columns)


# --------------------------------------------------------------------------
# Orquestração / CLI
# --------------------------------------------------------------------------

@dataclass
class ResultadoBatimento:
    planilhas_processadas: int
    total_verificado: int
    total_cadastrado: int
    total_pendente: int
    total_sem_numero: int
    caminho_saida: Path
    df_detalhe: pd.DataFrame = field(repr=False)
    df_resumo: pd.DataFrame = field(repr=False)


def executar_batimento(base_dir: Path, *, agora: Optional[datetime] = None) -> ResultadoBatimento:
    """Executa o fluxo completo de batimento a partir de `base_dir`
    (diretório que contém as pastas `entrada/` e `saida/`)."""
    estrutura = garantir_estrutura_pastas(base_dir)

    df_legal_one = carregar_base_legal_one(estrutura.base_legal_one)
    chaves_conhecidas = construir_chaves_conhecidas(df_legal_one)

    df_ged = carregar_export_ged(estrutura.export_ged)

    df_detalhe = processar_planilhas_cadastro(
        estrutura.planilhas_cadastro, chaves_conhecidas, df_ged
    )
    df_resumo = gerar_resumo(df_detalhe)

    agora = agora or datetime.now()
    nome_saida = f"Batimento_{agora.strftime('%Y%m%d_%H%M%S')}.xlsx"
    caminho_saida = estrutura.saida / nome_saida
    gerar_relatorio(df_detalhe, df_resumo, caminho_saida)

    contagem = df_detalhe["Status"].value_counts()
    return ResultadoBatimento(
        planilhas_processadas=df_detalhe["Arquivo de origem"].nunique(),
        total_verificado=len(df_detalhe),
        total_cadastrado=int(contagem.get(STATUS_CADASTRADO, 0)),
        total_pendente=int(contagem.get(STATUS_PENDENTE, 0)),
        total_sem_numero=int(contagem.get(STATUS_SEM_NUMERO, 0)),
        caminho_saida=caminho_saida,
        df_detalhe=df_detalhe,
        df_resumo=df_resumo,
    )


def main(argv: Optional[Sequence[str]] = None) -> int:
    base_dir = Path.cwd()
    try:
        resultado = executar_batimento(base_dir)
    except FileNotFoundError as erro:
        print(f"[ERRO] {erro}")
        return 1
    except ValueError as erro:
        print(f"[ERRO] {erro}")
        return 1

    print("=== Batimento de cadastro concluído ===")
    print(f"Planilhas processadas: {resultado.planilhas_processadas}")
    print(f"Processos verificados: {resultado.total_verificado}")
    print(f"Cadastrados: {resultado.total_cadastrado}")
    print(f"Pendentes: {resultado.total_pendente}")
    print(f"Sem número identificado: {resultado.total_sem_numero}")
    print(f"Relatório gerado em: {resultado.caminho_saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
